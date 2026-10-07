from pathlib import Path

import pdfplumber
from docx import Document

SUPPORTED_EXTENSIONS = {".pdf", ".docx"}
MIN_TEXT_LENGTH = 50


class UnsupportedFileError(Exception):
    pass


class EmptyResumeError(Exception):
    """Raised when no usable text is found (e.g. scanned/image-only PDF)."""
    pass


def extract_text_from_pdf(path: str | Path) -> str:
    pages = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            pages.append(page.extract_text() or "")
    return "\n".join(pages)


def extract_text_from_docx(path: str | Path) -> str:
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    # Many resumes put content inside tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def extract_text(path: str | Path) -> str:
    ext = Path(path).suffix.lower()
    if ext == ".pdf":
        text = extract_text_from_pdf(path)
    elif ext == ".docx":
        text = extract_text_from_docx(path)
    else:
        raise UnsupportedFileError(f"Unsupported file type: {ext}")

    if len(text.strip()) < MIN_TEXT_LENGTH:
        raise EmptyResumeError(
            "Could not extract text. The file may be a scanned image or empty."
        )
    return text